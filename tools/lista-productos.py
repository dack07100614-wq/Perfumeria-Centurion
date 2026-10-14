# Genera listas/Lista-de-productos.xlsx (precio, SKU, título, descripción, etiquetas) a partir de data/products.json
import json, os, re, csv, unicodedata
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
P = [p for p in json.load(open(ROOT + '/data/products.json', encoding='utf8'))['products'] if p.get('published', True)]
cod = {}
f = ROOT + '/codigos/Codigos-universales.csv'
if os.path.exists(f):
    for r in csv.reader(open(f, encoding='utf-8-sig')):
        if r and r[0] != 'id' and r[4]: cod[r[0]] = r[4]
def slug(s): return re.sub(r'[^A-Z0-9]+', '-', unicodedata.normalize('NFKD', s.upper()).encode('ascii', 'ignore').decode()).strip('-')
BR = {'Lattafa': 'LAT', 'Afnan': 'AFN', 'Armaf': 'ARM', 'Al Haramain': 'ALH', 'French Avenue': 'FRA', 'Rasasi': 'RAS', 'Maison Alhambra': 'MAL'}
def sku(p): return f"CEN-{BR.get(p['brand'], slug(p['brand'])[:3])}-{slug(p['name'])}-{p.get('ml', '')}"
def title(p):
    base = f"Perfume {p['brand']} {p['name']} {p.get('concentration','')} {p.get('ml','')}ml Original".replace('  ', ' ')
    if len(base) > 60: base = f"Perfume {p['brand']} {p['name']} {p.get('ml','')}ml Original"
    if len(base) > 60: base = f"{p['brand']} {p['name']} {p.get('ml','')}ml"
    return base[:60]
def desc(p):
    L = [f"{p['brand']} {p['name']} - {p.get('concentration','')} {p.get('ml','')} ml".replace('  ', ' '), '', p.get('description', '')]
    if p.get('family'): L.append(f"Familia olfativa: {p['family']}.")
    n = p.get('notes') or {}
    for k, lab in (('top', 'Notas de salida'), ('heart', 'Notas de corazón'), ('base', 'Notas de fondo')):
        if n.get(k): L.append(f"{lab}: {', '.join(n[k])}.")
    if p.get('longevity'): L.append(f"Duración aproximada: {p['longevity']}.")
    if p.get('projection'): L.append(f"Proyección: {p['projection']}.")
    if p.get('occasion'): L.append(f"Ideal para: {', '.join(p['occasion']).lower()}.")
    L += ['', 'Producto original.', 'Envíos a todo Uruguay.']
    return '\n'.join(L)
GEN = {'masculino': 'hombre', 'femenino': 'mujer', 'unisex': 'unisex'}
def tags(p):
    t = ['perfume árabe', 'perfume importado', p['brand'].lower(), p['name'].lower(), f"{p['brand'].lower()} {p['name'].lower()}"]
    t += [GEN.get(p.get('gender'), '')] + [x.lower() for x in p.get('tags', [])] + [x.lower() for x in p.get('keyNotes', [])[:3]]
    if p.get('concentration'): t.append(p['concentration'].lower())
    if p.get('ml'): t.append(f"{p['ml']} ml")
    seen = []; [seen.append(x) for x in t if x and x not in seen]
    return ', '.join(seen)
wb = Workbook(); ws = wb.active; ws.title = 'Productos'
H = ['SKU', 'Título', 'Precio (UYU)', 'Precio anterior', 'Descripción', 'Etiquetas', 'Marca', 'ml', 'Código universal (confirmar con la caja)', 'Foto']
ws.append(H)
for p in P:
    ws.append([sku(p), title(p), p['price'], p.get('oldPrice'), desc(p), tags(p), p['brand'], p.get('ml'), cod.get(p['id'], ''), p['id'] + '.jpg'])
for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='111111'); c.alignment = Alignment(vertical='center', wrap_text=True)
for i, w in enumerate([24, 46, 13, 13, 90, 60, 14, 7, 22, 30], 1): ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(vertical='top', wrap_text=True)
    row[2].number_format = '"$" #,##0'; row[3].number_format = '"$" #,##0'
ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
os.makedirs(ROOT + '/listas', exist_ok=True)
wb.save(ROOT + '/listas/Lista-de-productos.xlsx')
skus = [sku(p) for p in P]
print(len(P), 'productos;', len(set(skus)), 'SKU únicos;', max(len(title(p)) for p in P), 'caracteres máx. en títulos;', len(cod), 'con código')
