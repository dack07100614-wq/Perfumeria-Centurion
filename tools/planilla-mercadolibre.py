# Genera mercadolibre/Planilla-MercadoLibre.xlsx a partir de data/products.json
import json, os, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
P = [p for p in json.load(open(ROOT + '/data/products.json', encoding='utf8'))['products'] if p.get('published', True)]
photos = set(os.listdir(ROOT + '/mercadolibre/perfumes')) if os.path.isdir(ROOT + '/mercadolibre/perfumes') else set()
GEN = {'masculino': ('Hombre', 'Perfumes de Hombre'), 'femenino': ('Mujer', 'Perfumes de Mujer'), 'unisex': ('Unisex', 'Perfumes Unisex')}

def title(p):
    base = f"Perfume {p['brand']} {p['name']} {p.get('concentration','')} {p.get('ml','')}ml Original".replace('  ', ' ')
    if len(base) > 60: base = f"Perfume {p['brand']} {p['name']} {p.get('ml','')}ml Original"
    if len(base) > 60: base = f"{p['brand']} {p['name']} {p.get('ml','')}ml"
    return base[:60]

def desc(p):
    L = [f"{p['brand']} {p['name']} - {p.get('concentration','')} {p.get('ml','')} ml".replace('  ', ' '), '']
    L.append(p.get('description', ''))
    if p.get('family'): L.append(f"Familia olfativa: {p['family']}.")
    n = p.get('notes') or {}
    for k, lab in (('top', 'Notas de salida'), ('heart', 'Notas de corazón'), ('base', 'Notas de fondo')):
        if n.get(k): L.append(f"{lab}: {', '.join(n[k])}.")
    if p.get('longevity'): L.append(f"Duración aproximada: {p['longevity']}.")
    if p.get('projection'): L.append(f"Proyección: {p['projection']}.")
    if p.get('occasion'): L.append(f"Ideal para: {', '.join(p['occasion']).lower()}.")
    L += ['', 'Producto original.', 'Consultá por envíos a todo el país.']
    return '\n'.join(x for x in L if x is not None)

wb = Workbook(); ws = wb.active; ws.title = 'Publicaciones'
ws['A1'] = 'Comisión de Mercado Libre (completá el % real que te cobra):'; ws['A1'].font = Font(bold=True)
ws['E1'] = 0.15; ws['E1'].number_format = '0.0%'; ws['E1'].fill = PatternFill('solid', fgColor='FFF2CC'); ws['E1'].font = Font(bold=True)
ws['F1'] = '← valor de ejemplo, NO es el real. Cambialo.'; ws['F1'].font = Font(italic=True, color='C00000')
H = ['Título (máx. 60)', 'Marca', 'Perfume', 'Género', 'Categoría sugerida', 'Tamaño (ml)', 'Condición', 'Precio en tu página ($U)', 'Precio sugerido en Mercado Libre ($U)', 'Stock (completá)', 'Archivo de foto', 'Descripción']
for i, h in enumerate(H, 1):
    c = ws.cell(row=3, column=i, value=h); c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='111111'); c.alignment = Alignment(wrap_text=True, vertical='center')
thin = Side(style='thin', color='DDDDDD')
for r, p in enumerate(P, 4):
    g = GEN.get(p.get('gender'), GEN['unisex'])
    ph = f"{p['id']}.jpg" if f"{p['id']}.jpg" in photos else '(sin foto limpia: usar foto propia)'
    row = [title(p), p['brand'], p['name'], g[0], f"Belleza y Cuidado Personal > Perfumes > {g[1]}", p.get('ml'), 'Nuevo', p['price'], f'=ROUND(H{r}/(1-$E$1),-1)', None, ph, desc(p)]
    for i, v in enumerate(row, 1):
        c = ws.cell(row=r, column=i, value=v); c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = Border(bottom=thin)
    ws.cell(row=r, column=8).number_format = '"$" #,##0'; ws.cell(row=r, column=9).number_format = '"$" #,##0'
    ws.cell(row=r, column=10).fill = PatternFill('solid', fgColor='FFF2CC')
W = [44, 14, 28, 10, 34, 11, 11, 16, 20, 12, 30, 90]
for i, w in enumerate(W, 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[3].height = 42; ws.freeze_panes = 'A4'; ws.auto_filter.ref = f"A3:L{3+len(P)}"
ws2 = wb.create_sheet('Leeme')
for i, t in enumerate([
    'Cómo usar esta planilla',
    '1. Poné en la celda E1 de "Publicaciones" el porcentaje real de comisión que te cobra Mercado Libre. El 15% es solo un ejemplo.',
    '2. La columna "Precio sugerido" calcula el precio para que, descontada esa comisión, te quede el precio de tu página. No incluye el costo de envío.',
    '3. Completá el stock (celdas amarillas).',
    '4. Para cada perfume: copiá el título, la descripción y el precio al publicar. Las fotos están en la carpeta "perfumes" del ZIP, con el nombre de la columna "Archivo de foto".',
    '5. Elegí la categoría que te sugiera Mercado Libre al escribir el título. La de la planilla es orientativa.',
    '6. Revisá cada descripción antes de publicar: las notas olfativas y la duración son aproximadas.',
    '7. Mercado Libre puede pedirte comprobantes de originalidad de las marcas. Tené a mano las facturas de tu proveedor.'], 1):
    ws2.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True)
ws2['A1'].font = Font(bold=True, size=14); ws2.column_dimensions['A'].width = 120
os.makedirs(ROOT + '/mercadolibre', exist_ok=True)
wb.save(ROOT + '/mercadolibre/Planilla-MercadoLibre.xlsx')
print(len(P), 'perfumes;', max(len(title(p)) for p in P), 'caracteres como máximo en títulos')
